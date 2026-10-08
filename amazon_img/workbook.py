from __future__ import annotations

import re
import shutil
import os
import tempfile
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


def find_image_urls(path: Path, pattern: re.Pattern[str]) -> dict[str, list[tuple[str, str]]]:
    """Return unique matching URLs mapped to worksheet/cell coordinates."""
    keep_vba = path.suffix.lower() == ".xlsm"
    workbook = load_workbook(path, keep_vba=keep_vba, data_only=False, keep_links=True)
    found: dict[str, list[tuple[str, str]]] = defaultdict(list)
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        for match in pattern.finditer(cell.value):
                            # The configured expression may identify just a URL prefix.
                            # Extend it to the end of the URL while preserving surrounding text.
                            end = match.end()
                            while end < len(cell.value) and not cell.value[end].isspace() and cell.value[end] not in "\"'<>":
                                end += 1
                            url = cell.value[match.start():end].rstrip(",;)").rstrip(".")
                            if url.startswith(("http://", "https://")):
                                found[url].append((sheet.title, cell.coordinate))
    finally:
        workbook.close()
    return dict(found)


def make_copy_path(source: Path, task_id: str) -> Path:
    return source.with_name(f"{source.stem}.amazon-img-{task_id}{source.suffix}")


def create_workbook_copy(source: Path, destination: Path) -> None:
    shutil.copy2(source, destination)


class WorkbookEditor:
    """Keep one workbook open and atomically checkpoint modified cell values."""

    def __init__(self, path: Path):
        self.path = path
        self.workbook = load_workbook(
            path, keep_vba=path.suffix.lower() == ".xlsm", data_only=False, keep_links=True
        )

    def replace(self, sheet_name: str, coordinate: str, expected: str, replacement: str) -> bool:
        cell = self.workbook[sheet_name][coordinate]
        if isinstance(cell.value, str) and expected in cell.value:
            cell.value = cell.value.replace(expected, replacement)
            return True
        return False

    def save(self) -> None:
        fd, temp_name = tempfile.mkstemp(
            prefix=self.path.stem + ".", suffix=self.path.suffix, dir=self.path.parent
        )
        os.close(fd)
        try:
            self.workbook.save(temp_name)
            os.replace(temp_name, self.path)
        finally:
            if os.path.exists(temp_name):
                os.unlink(temp_name)

    def close(self) -> None:
        self.workbook.close()
